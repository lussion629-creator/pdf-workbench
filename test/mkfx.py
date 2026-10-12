import docx, pptx, openpyxl
from pptx.util import Inches, Pt
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4
from PIL import Image, ImageDraw, ImageFont
pdfmetrics.registerFont(TTFont('NG','/home/claude/pdfapp/fonts/NanumGothic.ttf'))
def mk(path, pages, title, change=False):
    c = canvas.Canvas(path, pagesize=A4)
    for i in range(pages):
        c.setFont('NG', 22); c.drawString(72, 770, f'{title} {i+1}쪽')
        c.setFont('NG', 11)
        y = 730
        lines = ['이 문서는 PDF 작업대 시험용 문서입니다.', '계약 금액은 1,250,000원이며 2026년 10월 31일까지 지급합니다.',
                 '담당자 연락처: 010-1234-5678, test@example.com', '주민등록번호 900101-1234567 은 가려야 합니다.',
                 'The quick brown fox jumps over the lazy dog.', '이름: ____________   서명: ____________', '□ 동의합니다   □ 동의하지 않습니다']
        if change and i == 0: lines[1] = '계약 금액은 1,500,000원이며 2026년 11월 30일까지 지급합니다.'
        for ln in lines: c.drawString(72, y, ln); y -= 22
        c.setFont('NG', 10); y -= 10
        for r, row in enumerate([['품목','수량','단가','금액'],['사과','10','1,000','10,000'],['배','5','2,500','12,500'],['합계','','','22,500']]):
            for k, v in enumerate(row): c.drawString(72 + k*110, y - r*18, v)
        c.setFillColorRGB(0.2,0.4,0.8); c.rect(72, 200, 200, 120, fill=1); c.setFillColorRGB(0,0,0)
        c.showPage()
    c.save()
mk('fx/doc3.pdf', 3, '시험 문서')
mk('fx/doc3b.pdf', 3, '시험 문서', change=True)
mk('fx/doc2.pdf', 2, '두번째 파일')
# scanned page (image only)
img = Image.new('RGB', (1240, 1754), 'white'); d = ImageDraw.Draw(img)
f = ImageFont.truetype('/home/claude/pdfapp/fonts/NanumGothic.ttf', 44)
for i, t in enumerate(['스캔 문서 인식 시험', '대한민국 서울특별시', 'Optical character recognition', '총 금액 35,000원']):
    d.text((120, 180 + i*110), t, fill='black', font=f)
img.save('fx/scan.jpg', quality=90)
c = canvas.Canvas('fx/scan.pdf', pagesize=A4); c.drawImage('fx/scan.jpg', 0, 0, *A4); c.showPage(); c.save()
img.resize((620, 877)).save('fx/photo.png')
# docx
dd = docx.Document(); dd.add_heading('워드 변환 시험', 1); dd.add_paragraph('첫 번째 문단입니다. 한글과 English가 섞여 있습니다. ' * 6)
dd.add_paragraph('목록 하나', style='List Bullet'); dd.add_paragraph('목록 둘', style='List Bullet')
t = dd.add_table(rows=2, cols=3); t.style = 'Table Grid'
for r in range(2):
    for k in range(3): t.cell(r,k).text = f'셀 {r}-{k}'
dd.add_picture('fx/photo.png', width=docx.shared.Inches(2)); dd.save('fx/test.docx')
# pptx
p = pptx.Presentation(); s = p.slides.add_slide(p.slide_layouts[0]); s.shapes.title.text = '발표 제목'; s.placeholders[1].text = '부제목 텍스트'
s2 = p.slides.add_slide(p.slide_layouts[1]); s2.shapes.title.text = '두 번째 슬라이드'; s2.placeholders[1].text_frame.text = '첫 줄 내용'; s2.placeholders[1].text_frame.add_paragraph().text = '둘째 줄 내용'
s2.shapes.add_picture('fx/photo.png', Inches(6), Inches(2), height=Inches(3))
p.save('fx/test.pptx')
# xlsx
wb = openpyxl.Workbook(); ws = wb.active; ws.title = '매출'
ws.append(['월', '매출', '비용']); [ws.append([f'{m}월', m*1000, m*600]) for m in range(1, 13)]
wb.create_sheet('메모').append(['비고', '시험'])
wb.save('fx/test.xlsx')
open('fx/page.html','w').write('<html><head><style>h1{color:#2747d4}</style></head><body><h1>HTML 변환 시험</h1><p>문단입니다.</p><table border=1><tr><td>a</td><td>b</td></tr></table></body></html>')
print('fixtures ok')
