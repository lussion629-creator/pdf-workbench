import os, sys, subprocess, zipfile, json, time, threading, http.server, functools, traceback
from playwright.sync_api import sync_playwright

ROOT = os.path.dirname(os.path.abspath(__file__))
DIST = os.path.join(ROOT, '..', 'dist', sys.argv[1] if len(sys.argv) > 1 else 'app')
FX = os.path.join(ROOT, 'fx'); OUT = os.path.join(ROOT, 'out'); os.makedirs(OUT, exist_ok=True)
ONLY = sys.argv[2].split(',') if len(sys.argv) > 2 else None

class Q(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a): pass
srv = http.server.ThreadingHTTPServer(('127.0.0.1', 8765), functools.partial(Q, directory=DIST))
threading.Thread(target=srv.serve_forever, daemon=True).start()
BASE = 'http://127.0.0.1:8765/index.html'

def pdf_text(p):
    return subprocess.run(['pdftotext', '-layout', p, '-'], capture_output=True, text=True).stdout
def check_file(p):
    ext = p.rsplit('.', 1)[-1].lower()
    if ext == 'pdf':
        r = subprocess.run(['qpdf', '--check'] + (['--password=abc123'] if 'protect' in p else []) + [p], capture_output=True, text=True)
        info = subprocess.run(['pdfinfo', p], capture_output=True, text=True).stdout
        pages = [l for l in info.splitlines() if l.startswith('Pages:')]
        return (r.returncode in (0, 3)), (pages[0] if pages else info[:80]) + ('' if r.returncode == 0 else ' qpdf rc=%d %s' % (r.returncode, (r.stdout + r.stderr)[-200:]))
    if ext in ('docx', 'pptx', 'xlsx', 'zip'):
        z = zipfile.ZipFile(p); bad = z.testzip()
        if ext == 'docx':
            import docx; d = docx.Document(p); return bad is None, f'{len(d.paragraphs)} paragraphs: ' + ' | '.join(x.text for x in d.paragraphs[:4])[:160]
        if ext == 'pptx':
            import pptx; d = pptx.Presentation(p); return bad is None, f'{len(d.slides)} slides'
        if ext == 'xlsx':
            import openpyxl; wb = openpyxl.load_workbook(p); ws = wb.worksheets[0]; return bad is None, f'sheets={wb.sheetnames} first rows=' + str([[c.value for c in r] for r in ws.iter_rows(max_row=3)])[:200]
        return bad is None, f'{len(z.namelist())} entries'
    if ext in ('jpg', 'png'):
        from PIL import Image; im = Image.open(p); return True, f'{im.size}'
    return os.path.getsize(p) > 0, open(p, encoding='utf-8', errors='replace').read(160).replace('\n', ' ')

class R(list):
    def append(self, x):
        super().append(x); print(('PASS' if x[1] else 'FAIL'), x[0], '-', x[2], flush=True)
        [print('     ', 'ok ' if o[1] else 'BAD', o[0], '-', o[2], flush=True) for o in x[3]]
results = R()
def run(page, tid, files, actions=None, expect_outputs=True, timeout=180000, run_click=True):
    if ONLY and tid not in ONLY: return
    t0 = time.time(); errs = []
    hnd = lambda e: errs.append(str(e)); page.on('pageerror', hnd)
    try:
        page.goto('about:blank'); page.goto(BASE + '#' + tid); page.wait_for_selector('.tool-head')
        if files:
            with page.expect_file_chooser() as fc: page.get_by_role('button', name='파일 선택').click()
            fc.value.set_files([os.path.join(FX, f) for f in files])
            page.wait_for_timeout(900)
        if actions: actions(page)
        if not run_click: return
        page.click('#runbtn', timeout=10000)
        page.wait_for_selector('.done-head, .rail .err', timeout=timeout)
        if page.query_selector('.rail .err'):
            results.append((tid, False, 'ERROR: ' + page.inner_text('.rail .err'), [])); return
        outs = []
        for i, b in enumerate(page.query_selector_all('.out .btn')):
            with page.expect_download() as dl: b.click()
            d = dl.value; p = os.path.join(OUT, f'{tid}_{len(results)}_{i}_' + d.suggested_filename); d.save_as(p)
            ok, info = check_file(p); outs.append((os.path.basename(p), ok, info))
        ok = all(o[1] for o in outs) and (bool(outs) or not expect_outputs)
        head = page.inner_text('.done-head')
        results.append((tid, ok, head.replace('\n', ' / ') + f' [{time.time()-t0:.1f}s]' + (' JSERR:' + ';'.join(errs) if errs else ''), outs))
    except Exception as e:
        results.append((tid, False, 'EXC ' + str(e)[:300] + (' JSERR:' + ';'.join(errs) if errs else ''), []))
    finally:
        page.remove_listener('pageerror', hnd)

def seg(page, label): page.locator('.rail .seg button', has_text=label).first.click()

with sync_playwright() as pw:
    br = pw.chromium.launch(args=['--no-sandbox'])
    ctx = br.new_context(accept_downloads=True, viewport={'width': 1280, 'height': 900})
    page = ctx.new_page()
    page.on('console', lambda m: m.type == 'error' and print('  console:', m.text[:200]))
    run(page, 'merge', ['doc3.pdf', 'doc2.pdf'])
    run(page, 'split', ['doc3.pdf'])
    run(page, 'split', ['doc3.pdf'], lambda p: seg(p, '모든 페이지'))
    run(page, 'remove', ['doc3.pdf'], lambda p: (p.fill('#o-rmRange', '2'), p.dispatch_event('#o-rmRange', 'change'), p.wait_for_timeout(300)))
    run(page, 'extract', ['doc3.pdf'], lambda p: (p.locator('.pcard').nth(0).click(), p.locator('.pcard').nth(2).click()))
    run(page, 'organize', ['doc3.pdf'], lambda p: (p.locator('.pcard').nth(1).hover(), p.locator('.pcard').nth(1).get_by_label('삭제').click(force=True)))
    run(page, 'scan', ['scan.jpg', 'photo.png'])
    run(page, 'compress', ['doc3.pdf', 'scan.pdf'])
    run(page, 'repair', ['doc3.pdf'])
    run(page, 'ocr', ['scan.pdf'], timeout=300000)
    run(page, 'jpg2pdf', ['scan.jpg', 'photo.png'])
    run(page, 'word2pdf', ['test.docx'])
    run(page, 'word2pdf', ['test.docx'], lambda p: seg(p, '모양 그대로'))
    run(page, 'ppt2pdf', ['test.pptx'])
    run(page, 'excel2pdf', ['test.xlsx'])
    run(page, 'html2pdf', ['page.html'])
    run(page, 'pdf2jpg', ['doc2.pdf'])
    run(page, 'pdf2jpg', ['scan.pdf'], lambda p: seg(p, '들어있는 이미지 추출'))
    run(page, 'pdf2word', ['doc3.pdf'])
    run(page, 'pdf2ppt', ['doc2.pdf'])
    run(page, 'pdf2excel', ['doc2.pdf'])
    run(page, 'pdf2pdfa', ['doc2.pdf'])
    run(page, 'rotate', ['doc3.pdf'], lambda p: p.locator('.pcard').nth(0).click())
    run(page, 'pagenum', ['doc3.pdf'], lambda p: p.select_option('#o-fmt', '{n} / {p}'))
    run(page, 'watermark', ['doc3.pdf'], lambda p: p.locator('.rail label.chk', has_text='바둑판').click())
    run(page, 'watermark', ['doc3.pdf'], lambda p: seg(p, '내용 아래'))
    def crop_act(p): p.wait_for_selector('.cropbox', timeout=15000)
    run(page, 'crop', ['doc3.pdf'], crop_act)
    def edit_act(p):
        p.wait_for_selector('.ed-layer', timeout=15000)
        p.locator('.ed-bar .tb', has_text='글자').click(); box = p.locator('.ed-layer').bounding_box()
        p.mouse.click(box['x'] + 100, box['y'] + 120); p.keyboard.type('추가한 글자 ABC'); p.mouse.click(box['x'] + 400, box['y'] + 700)
        p.locator('.ed-bar .tb', has_text='사각형').click(); p.mouse.move(box['x'] + 100, box['y'] + 300); p.mouse.down(); p.mouse.move(box['x'] + 300, box['y'] + 400); p.mouse.up()
        p.get_by_role('button', name='펜', exact=True).click(); p.mouse.move(box['x'] + 120, box['y'] + 500); p.mouse.down();
        for k in range(10): p.mouse.move(box['x'] + 120 + k * 15, box['y'] + 500 + (k % 3) * 10)
        p.mouse.up(); p.wait_for_timeout(300)
    run(page, 'edit', ['doc3.pdf'], edit_act)
    def form_act(p):
        p.wait_for_selector('.ed-layer', timeout=15000); p.get_by_role('button', name='빈칸 자동 찾기').click(); p.wait_for_timeout(1500)
    run(page, 'forms', ['doc3.pdf'], form_act)
    run(page, 'unlock', ['locked.pdf'], lambda p: (p.fill('#o-pw', 'user'), p.dispatch_event('#o-pw', 'change')))
    run(page, 'protect', ['doc2.pdf'], lambda p: (p.fill('#o-pw', 'abc123'), p.dispatch_event('#o-pw', 'change'), p.fill('#o-pw2', 'abc123'), p.dispatch_event('#o-pw2', 'change')))
    def sign_act(p):
        p.wait_for_selector('.ed-layer', timeout=15000); p.locator('.ed-bar .tb', has_text='서명').first.click()
        p.get_by_role('button', name='입력하기').click(); p.fill('#sig-name', '홍길동'); p.get_by_role('button', name='서명 넣기').click(); p.wait_for_timeout(800)
    run(page, 'sign', ['doc2.pdf'], sign_act)
    def red_act(p):
        p.wait_for_selector('.ed-layer', timeout=15000)
        p.locator('.rail label.chk', has_text='주민등록번호').click(); p.locator('.rail label.chk', has_text='전화번호').click()
        p.get_by_role('button', name='찾아서 표시하기').click(); p.wait_for_timeout(1500)
    run(page, 'redact', ['doc2.pdf'], red_act)
    run(page, 'compare', ['doc3.pdf', 'doc3b.pdf'], expect_outputs=True)
    run(page, 'compare', ['doc3.pdf', 'doc3b.pdf'], lambda p: seg(p, '모양 비교'), expect_outputs=False)
    run(page, 'markdown', ['doc3.pdf'])
    # home page renders all tools
    page.goto(BASE); page.wait_for_selector('.card'); n = len(page.query_selector_all('.card'))
    results.append(('home', n >= 33, f'{n} tool cards', []))
    br.close()

okc = 0
for tid, ok, msg, outs in results:
    okc += ok
    print(('PASS' if ok else 'FAIL'), tid, '-', msg)
    for o in outs: print('     ', 'ok ' if o[1] else 'BAD', o[0], '-', o[2])
print(f'{okc}/{len(results)} passed')
